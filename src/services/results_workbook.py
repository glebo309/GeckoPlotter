"""The results of a workspace as one Excel workbook to keep working with: a summary per compound and sample, every
peak, the calibrations, the samples, the analysis history and what made the file. Plain tables with a header row,
filters and frozen headers; numbers stay numbers. Used by the Export panel and the agent (gecko_api.export_results).
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import List

HEADER_FILL = '2C5240'
BAND_FILL = 'F3F6F4'


def _channel(bundle, component) -> str:
    return bundle.component_display_name(component)


def peak_rows(store) -> List[dict]:
    """One row per peak: sample, channel, name, retention time, bounds, area, height, area %, concentration."""
    from services.natural_sort import natural_key
    rows = []
    bundles = sorted(store.iter_bundles(), key=lambda b: natural_key(b.name))
    for bundle in bundles:
        for component in bundle.component_list:
            peaks = sorted(store.peak_service.get_peaks(component.component_id) or [], key=lambda p: p.rt)
            total = sum(max(float(p.area or 0), 0) for p in peaks)
            for p in peaks:
                rows.append({
                    'Sample': bundle.name, 'Channel': _channel(bundle, component),
                    'Peak': p.label or p.name, 'RT (min)': round(float(p.rt), 4),
                    'Start (min)': round(float(p.rt_start), 4), 'End (min)': round(float(p.rt_end), 4),
                    'Area': float(p.area), 'Height': float(p.height),
                    'Area %': round(100 * float(p.area) / total, 2) if total > 0 else None,
                    'Width (min)': round(float(p.width), 4) if p.width else None,
                    'Concentration': p.concentration, 'Integration': p.source})
    return rows


def _summary(rows, value):
    """Compound (peak name, channel) x sample table of ``value``, samples in the tree's natural order."""
    from services.natural_sort import natural_key
    samples = sorted({r['Sample'] for r in rows}, key=natural_key)
    compounds = []
    for r in rows:
        key = (r['Peak'], r['Channel'])
        if key not in compounds:
            compounds.append(key)
    compounds.sort(key=lambda k: (k[1], min(r['RT (min)'] for r in rows if (r['Peak'], r['Channel']) == k)))
    table = []
    for name, channel in compounds:
        line = {'Compound': name, 'Channel': channel}
        for sample in samples:
            hits = [r[value] for r in rows if r['Sample'] == sample and r['Peak'] == name and r['Channel'] == channel
                    and r[value] is not None]
            line[sample] = sum(hits) if hits else None
        table.append(line)
    return table, ['Compound', 'Channel'] + samples


def _write(sheet, header, rows, widths=None, note=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    start = 1
    if note:
        sheet.cell(row=1, column=1, value=note).font = Font(italic=True, color='6E7378')
        start = 3
    for col, name in enumerate(header, start=1):
        cell = sheet.cell(row=start, column=col, value=name)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for r, row in enumerate(rows, start=start + 1):
        for col, name in enumerate(header, start=1):
            value = row.get(name)
            cell = sheet.cell(row=r, column=col, value=value)
            if isinstance(value, float):
                cell.number_format = '0.00' if abs(value) < 1e5 else '0.000E+00'
            if (r - start) % 2 == 0:
                cell.fill = PatternFill('solid', fgColor=BAND_FILL)
    sheet.freeze_panes = sheet.cell(row=start + 1, column=1)
    if rows:
        sheet.auto_filter.ref = f'A{start}:{get_column_letter(len(header))}{start + len(rows)}'
    for col, name in enumerate(header, start=1):
        width = (widths or {}).get(name) or min(max(len(str(name)), *(len(str(row.get(name) or ''))
                                                                     for row in rows[:200] or [{}])) + 2, 40)
        sheet.column_dimensions[get_column_letter(col)].width = max(width, 9)


def build_results_workbook(store, project_name: str = '') -> bytes:
    """The workbook as .xlsx bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from services import about
    from services.research_package import agent_note

    rows = peak_rows(store)
    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    table, header = _summary(rows, 'Area')
    _write(summary, header, table, note='Peak areas per compound and sample (sum if a name occurs twice).')
    if any(r['Concentration'] is not None for r in rows):
        conc, conc_header = _summary(rows, 'Concentration')
        _write(book.create_sheet('Concentrations'), conc_header, conc, note='Concentrations from the calibrations.')

    peak_header = ['Sample', 'Channel', 'Peak', 'RT (min)', 'Start (min)', 'End (min)', 'Area', 'Height', 'Area %',
                   'Width (min)', 'Concentration', 'Integration']
    _write(book.create_sheet('Peaks'), peak_header, rows)

    calibrations = []
    for model in store.calibration_service.get_all_models().values():
        coefficients = list(model.coefficients or [])
        for point in model.standards:
            calibrations.append({'Analyte': model.analyte_name, 'Fit': model.fit_type, 'Through zero': model.force_zero,
                                 'Coefficients': ', '.join(f'{c:.6g}' for c in coefficients),
                                 'R²': model.r_squared, 'Standard': point.sample_name,
                                 'Concentration': point.concentration, 'Units': model.units,
                                 'Response': point.area, 'Used': point.included})
    if calibrations:
        _write(book.create_sheet('Calibration'), ['Analyte', 'Fit', 'Through zero', 'Coefficients', 'R²', 'Standard',
                                                 'Concentration', 'Units', 'Response', 'Used'], calibrations)

    from services.natural_sort import natural_key
    samples = []
    for bundle in sorted(store.iter_bundles(), key=lambda b: natural_key(b.name)):
        meta = (bundle.primary.metadata if bundle.primary else {}) or {}
        samples.append({'Sample': bundle.name,
                        'Channels': ', '.join(_channel(bundle, c) for c in bundle.component_list),
                        'Files': ', '.join(dict.fromkeys(c.original_filename for c in bundle.component_list)),
                        'Acquired': str(meta.get('acquired_at') or meta.get('acquisition_date') or meta.get('date') or '')
                        or None,
                        'Method': str(meta.get('method_name') or meta.get('method') or '') or None})
    _write(book.create_sheet('Samples'), ['Sample', 'Channels', 'Files', 'Acquired', 'Method'], samples,
           widths={'Files': 60, 'Channels': 40})

    history = []
    for entry in getattr(store, 'analysis_history', None) or []:
        agent = entry.get('agent') or {}
        history.append({'Time (UTC)': str(entry.get('at', '')).replace('T', ' ')[:19],
                        'Change': '; '.join(entry.get('trigger') or []) or ', '.join(sorted(entry.get('changes', {}))),
                        'By': ('MCP agent: ' + ' '.join(str(x) for x in (agent.get('client'), agent.get('model')) if x))
                        if agent else 'GeckoPlotter app'})
    if history:
        _write(book.create_sheet('History'), ['Time (UTC)', 'Change', 'By'], history, widths={'Change': 70, 'By': 36})

    info = book.create_sheet('About')
    lines = [('Project', project_name or (getattr(store, 'session_meta', {}) or {}).get('project_name') or 'Untitled'),
             ('Exported', datetime.now().strftime('%d %b %Y, %H:%M')),
             ('GeckoPlotter version', about.VERSION.get('sha', '')[:7] or 'unknown'),
             ('Agent', agent_note(store)),
             ('Peaks', str(len(rows))), ('Samples', str(len(samples)))]
    for r, (key, value) in enumerate(lines, start=1):
        info.cell(row=r, column=1, value=key).font = Font(bold=True)
        info.cell(row=r, column=2, value=value)
    info.column_dimensions['A'].width = 22
    info.column_dimensions['B'].width = 90

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
