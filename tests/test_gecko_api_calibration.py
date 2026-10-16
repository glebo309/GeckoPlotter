"""The agent calibrates a peak from standards as the app's calibration editor does, and quantifies the other samples."""
import os
import sys

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'src'))

from gecko_api import GeckoError, Workspace  # noqa: E402
from services.data_store import Bundle, BundleType, Component, ComponentRole  # noqa: E402
from src.ingestion.types import ColumnRole, SourceCategory  # noqa: E402

STANDARDS = {'Ala 0.31 mM': 0.3125, 'Ala 0.62 mM': 0.625, 'Ala 1.25 mM': 1.25, 'Ala 2.5 mM': 2.5, 'Ala 5 mM': 5.0,
             'Ala 10 mM': 10.0}


def _sample(name, amount, rt=7.3):
    t = np.linspace(0, 12, 2400)
    y = 40.0 * amount * np.exp(-((t - rt) / 0.06) ** 2) + 0.2 * np.sin(t * 7)
    frame = pd.DataFrame({'Time (min)': t, 'Absorbance (mAU)': y})
    meta = {'wavelength_nm': 210, 'auto_x': 'Time (min)', 'auto_y': ['Absorbance (mAU)'], 'sample_name': name}
    component = Component(name + '-uv', ComponentRole.PRIMARY, name + '.txt', dataframe=frame,
                          column_roles={'Time (min)': ColumnRole.X, 'Absorbance (mAU)': ColumnRole.Y},
                          source_category=SourceCategory.CHROMATOGRAM, metadata=meta)
    return Bundle(name, name, BundleType.CHROMATOGRAM, component_list=[component])


@pytest.fixture
def workspace():
    ws = Workspace()
    for name, amount in {**STANDARDS, 'Unknown A': 1.5, 'Unknown B': 20.0}.items():
        ws.store.add_bundle(_sample(name, amount))
    ws.store.add_bundle(_sample('Shifted', 2.0, rt=7.8))
    ws.integrate(7.0, 7.6, samples=[*STANDARDS, 'Unknown A', 'Unknown B'], name='Ala')
    ws.integrate(7.5, 8.1, samples='Shifted', name='Ala')
    ws.store.add_bundle(_sample('Empty', 0.0))              # a standard that lacks the peak
    return ws


def test_a_linear_calibration_quantifies_the_unknowns_and_flags_extrapolation(workspace):
    result = workspace.calibrate('Ala', STANDARDS, units='mM')
    assert result['r_squared'] > 0.999 and set(result['coefficients']) == {'slope', 'intercept'}
    by_sample = {c['sample']: c for c in result['concentrations']}
    assert by_sample['Unknown A']['concentration'] == pytest.approx(1.5, rel=0.01)
    assert by_sample['Unknown B']['status'] == 'Outside calibration range'
    assert any('Unknown B' in w for w in result['warnings'])
    assert any('Shifted' in w and 'window' in w for w in result['warnings'])
    rows = {r['sample']: r for r in workspace.peaks() if r['name'] == 'Ala'}
    assert rows['Unknown A']['concentration'] == pytest.approx(1.5, rel=0.01)
    assert any(e['trigger'][0].startswith('Calibrated Ala') for e in workspace.store.analysis_history)


def test_a_missing_standard_is_reported_and_too_few_standards_refused(workspace):
    result = workspace.calibrate('Ala', {**STANDARDS, 'Empty': 0.0})
    assert result['r_squared'] > 0.999 and any('Empty has no peak' in w for w in result['warnings'])
    with pytest.raises(GeckoError, match='at least two standards'):
        workspace.calibrate('Nope', STANDARDS)


def test_recalibrating_replaces_the_curve_and_removing_clears_concentrations(workspace):
    workspace.calibrate('Ala', STANDARDS)
    again = workspace.calibrate('Ala', STANDARDS, fit='quadratic')
    assert set(again['coefficients']) == {'a', 'b', 'c'} and len(workspace.list_calibrations()) == 1
    assert workspace.remove_calibration('Ala') == []
    assert all(r['concentration'] is None for r in workspace.peaks())


def test_the_calibration_is_saved_and_in_the_workbook(workspace, tmp_path):
    from openpyxl import load_workbook
    workspace.calibrate('Ala', STANDARDS)
    again = Workspace()
    again.open(workspace.save(str(tmp_path / 'cal')))
    assert again.list_calibrations()[0]['calibration'] == 'Ala'
    assert {r['sample']: r['concentration'] for r in again.peaks()}['Unknown A'] == pytest.approx(1.5, rel=0.01)
    book = load_workbook(again.export_results(str(tmp_path / 'cal')))
    assert 'Calibration' in book.sheetnames


def test_calibration_works_on_samples_read_from_files(tmp_path, monkeypatch):
    """One channel per sample (CSV files): the editor files peaks under the sample, not the channel."""
    t = np.arange(0, 6, 0.01)
    levels = {'Std 1': 1.0, 'Std 2': 2.0, 'Std 4': 4.0, 'Unknown': 3.0}
    for name, c in levels.items():
        y = 50 * c * np.exp(-((t - 3.0) / 0.05) ** 2)
        (tmp_path / f'{name}.csv').write_text('Time (min),Absorbance (mAU)\n' +
                                               '\n'.join(f'{a:.2f},{b:.4f}' for a, b in zip(t, y)))
    # In a full run tests/ingestion is imported as ``ingestion`` first; reading files needs the app's package.
    saved = {k: v for k, v in sys.modules.items() if k == 'ingestion' or k.startswith('ingestion.')}
    for name in saved:
        monkeypatch.delitem(sys.modules, name)
    monkeypatch.syspath_prepend(os.path.join(os.path.dirname(__file__), '..', 'src'))
    ws = Workspace()
    ws.open(str(tmp_path))
    for name in [k for k in sys.modules if k == 'ingestion' or k.startswith('ingestion.')]:
        if name not in saved:
            del sys.modules[name]
    ws.integrate(2.8, 3.2, name='Product')
    result = ws.calibrate('Product', {k: v for k, v in levels.items() if k.startswith('Std')})
    assert {c['sample']: c['concentration'] for c in result['concentrations']}['Unknown'] == pytest.approx(3.0, rel=0.01)


def test_other_standard_peaks_are_not_unknowns_and_named_unknown_wins(workspace, tmp_path):
    from copy import deepcopy
    for name in [*STANDARDS, 'Unknown A']:
        peaks = workspace.store.peak_service.get_peaks(name + '-uv')
        other = deepcopy(peaks[0])
        other.peak_id += '-other'
        other.name = other.label = 'main peak'
        other.rt += .03
        owner = workspace.store.peak_owner_id(name + '-uv')
        workspace.store.peak_service._peaks[owner].append(other)
    result = workspace.calibrate('Ala', STANDARDS, name='Alanine curve')
    assert {r['sample'] for r in result['concentrations']} == {'Unknown A', 'Unknown B'}
    assert all(r['concentration'] is not None for r in result['concentrations'])
    assert not any('Ambiguous' in w for w in result['warnings'])
    # The same shared matching is used when a project is reopened and refreshed.
    again = Workspace()
    again.open(workspace.save(str(tmp_path / 'named')))
    again._refresh_concentrations()
    assert all(r['concentration'] is None for r in again.peaks() if r['name'] == 'main peak')
    assert next(r for r in again.peaks('Unknown A') if r['name'] == 'Ala')['concentration'] == pytest.approx(1.5, rel=.01)


def test_below_intercept_standard_explains_missing_back_calculation(workspace):
    for name, area in zip(STANDARDS, [0.01, 0.08, 0.28, 0.36, 1.26, 0.39]):
        workspace.store.peak_service.get_peaks(name + '-uv')[0].area = area
    result = workspace.calibrate('Ala', STANDARDS)
    for row in result['standards'][:2]:
        assert row['back_calculated'] is None
        assert row['note'] == 'area below the intercept'
