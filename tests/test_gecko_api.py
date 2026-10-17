"""The agent interface (src/gecko_api.py, gecko_mcp.py) on a small LC-MS sequence written here: two run folders, each
with full-scan spectra (ANDI netCDF) and a UV trace, as Agilent exports them."""
import os
import sys

import numpy as np
import pytest

netCDF4 = pytest.importorskip('netCDF4')

ROOT = os.path.join(os.path.dirname(__file__), '..')
sys.path.insert(0, os.path.join(ROOT, 'src'))

from gecko_api import GeckoError, Workspace  # noqa: E402

TIMES = np.linspace(0.01, 14.0, 490)
ION = 633.8                       # the compound, at 8.4 min
BACKGROUND = 60.1


def _spectra(path, amount):
    """Full-scan positive spectra: the compound (with its 13C isotope) eluting at 8.4 min over a constant ion."""
    profile = amount * 6e5 * np.exp(-((TIMES - 8.4) / 0.06) ** 2)
    mz, intensity, index, counts = [], [], [], []
    for height in profile:
        index.append(len(mz))
        ions = [(BACKGROUND, 5e4), (ION, height), (ION + 1, 0.45 * height)]
        mz += [m for m, _ in ions]
        intensity += [v for _, v in ions]
        counts.append(len(ions))
    with netCDF4.Dataset(path, 'w') as ds:
        ds.createDimension('scan_number', len(TIMES))
        ds.createDimension('point_number', len(mz))
        for name, values, kind, dim in (
                ('scan_acquisition_time', TIMES * 60, 'f8', 'scan_number'),
                ('total_intensity', [sum(intensity[i:i + n]) for i, n in zip(index, counts)], 'f8', 'scan_number'),
                ('scan_index', index, 'i4', 'scan_number'), ('point_count', counts, 'i4', 'scan_number'),
                ('mass_values', mz, 'f4', 'point_number'), ('intensity_values', intensity, 'f4', 'point_number')):
            ds.createVariable(name, kind, (dim,))[:] = np.asarray(values)
        ds.experiment_title = os.path.basename(path).split('_MS1')[0] + '.dx'


def _run(folder, stamp, amount):
    folder.mkdir()
    base = f'2026-09-10 {stamp}+02-00-07'
    _spectra(str(folder / f'{base}_MS1 +TIC SCAN ESI Frag=135V Gain=1-0_spectra.cdf'), amount)
    uv = amount * 800 * np.exp(-((TIMES - 7.3) / 0.05) ** 2)
    (folder / f'{base}.dx_VWD1A.CSV').write_text('\n'.join(f'{t:.4f},{v:.4f}' for t, v in zip(TIMES, uv)) + '\n')


@pytest.fixture(autouse=True)
def real_ingestion(monkeypatch):
    """In a full run ``tests/ingestion`` is imported as ``ingestion`` first; these tests read files with the app's."""
    def ours():
        return {k: v for k, v in sys.modules.items() if k == 'ingestion' or k.startswith('ingestion.')}
    before = ours()
    for name in before:
        del sys.modules[name]
    monkeypatch.syspath_prepend(os.path.join(ROOT, 'src'))
    yield
    for name in ours():
        del sys.modules[name]
    sys.modules.update(before)


@pytest.fixture
def sequence(tmp_path, monkeypatch):
    monkeypatch.setenv('GECKO_SETTINGS_DIR', str(tmp_path / 'settings'))
    root = tmp_path / 'HalB'
    root.mkdir()
    _run(root / '2026-10-15 11-05-28+02-00-4', '15-53-54', 2.0)
    _run(root / '2026-10-15 11-05-28+02-00-CONTROL 4', '15-39-09', 1.0)
    return root


def test_a_folder_of_run_folders_opens_as_one_named_sample_per_run(sequence):
    ws = Workspace()
    ws.open(str(sequence))
    names = sorted(s['name'] for s in ws.samples() if s['type'] == 'lcms')
    assert names == ['4', 'CONTROL 4']
    channels = {c['name']: c for c in next(s for s in ws.samples() if s['name'] == '4')['channels']}
    assert channels['+TIC']['mass_spectra'] and channels['+TIC']['kind'] == 'ms+'
    with pytest.raises(GeckoError, match='Loaded: '):
        ws._sample('nope')


def test_eic_areas_follow_the_amount_and_the_spectrum_names_the_ion(sequence):
    ws = Workspace()
    ws.open(str(sequence))
    rows = {r['sample']: r for r in ws.eic(ION, 0.5)}
    assert rows['4']['apex_rt'] == pytest.approx(8.4, abs=0.03)
    assert rows['4']['area'] / rows['CONTROL 4']['area'] == pytest.approx(2.0, rel=0.02)
    spectrum = ws.mass_spectrum('4', rt_start=8.3, rt_end=8.5, top=3)
    assert [p['mz'] for p in spectrum['peaks']][:2] == [ION, round(ION + 1, 2)]


def test_an_integration_is_saved_in_a_project_that_reopens_with_the_peak(sequence, tmp_path):
    ws = Workspace()
    ws.open(str(sequence))
    peak = ws.integrate(8.2, 8.6, samples='4', channel='+TIC', name='m634 compound')[0]
    assert peak['name'] == 'm634 compound' and peak['area'] > 0
    assert ws.mass_spectrum('4', peak='m634 compound', top=1)['peaks'][0]['mz'] == ION
    project = ws.save(str(tmp_path / 'halb'))
    again = Workspace()
    again.open(project)
    reopened = [r for r in again.peaks() if r['name'] == 'm634 compound']
    assert len(reopened) == 1 and reopened[0]['area'] == pytest.approx(peak['area'])
    shown = {b.name for b in again._bundles() if again.store.visibility_state.get(b.bundle_id, True)}
    assert shown == {'4'}                                    # it opens on the run that holds the result
    from services import session_service
    report = session_service.restore_session(again.store, project)
    assert report.client_state['active-file'] == again._channel(again._sample('4'), '+TIC').component_id


def test_detected_peaks_are_the_detect_buttons_peaks(sequence):
    from services.detection import detect
    ws = Workspace()
    ws.open(str(sequence))
    rows = ws.detect_peaks(samples='4', channel='UV', sensitivity=50)
    assert [round(r['rt'], 1) for r in rows] == [7.3]
    component = ws._channel(ws._sample('4'), 'UV')
    x, y = ws._xy(component.component_id)
    again = detect(ws.store, component.component_id, x, y, 50)
    assert [p.area for p in again] == [pytest.approx(rows[0]['area'], rel=1e-5)]


def test_the_mcp_server_offers_the_tools():
    pytest.importorskip('mcp')
    sys.path.insert(0, ROOT)
    import gecko_mcp
    import asyncio
    names = {tool.name for tool in asyncio.run(gecko_mcp.server.list_tools())}
    assert {'open_data', 'list_samples', 'detect_peaks', 'integrate', 'peak_table', 'extract_ion_chromatogram',
            'mass_spectrum', 'figure', 'save_project', 'export_results', 'calibrate', 'list_calibrations',
            'remove_calibration', 'remove_peak', 'identify_agent', 'analysis_guide'} <= names
    assert 'analysis_guide' in gecko_mcp.server.instructions


def test_an_integration_window_on_the_flanks_takes_the_whole_peak(sequence):
    ws = Workspace()
    ws.open(str(sequence))
    whole = ws.integrate(7.9, 8.9, samples='CONTROL 4', channel='+TIC', snap_to_valleys=False)[0]
    flanks = ws.integrate(8.35, 8.45, samples='4', channel='+TIC')[0]
    assert flanks['rt_start'] < 8.3 and flanks['rt_end'] > 8.5
    assert flanks['area'] == pytest.approx(2 * whole['area'], rel=0.03)      # sample 4 holds twice the amount


def test_duplicate_agilent_exports_are_left_out_when_the_spectra_are_there():
    from services.workspace_loader import split_duplicate_exports
    files = [(n, b'') for n in ('run_MS1 +TIC SCAN ESI_spectra.cdf', 'run.dx_VWD1A.CSV', '4.cdf', '4_004.cdf',
                                'run.dx_MS1 +EIC(633.7-634.7) ESI.CSV', 'method.lcm')]
    kept, duplicates = split_duplicate_exports(files)
    assert [n for n, _ in kept] == ['run_MS1 +TIC SCAN ESI_spectra.cdf', 'run.dx_VWD1A.CSV', 'method.lcm']
    assert [n for n, _ in duplicates] == ['4.cdf', '4_004.cdf', 'run.dx_MS1 +EIC(633.7-634.7) ESI.CSV']
    assert split_duplicate_exports(files[1:]) == (files[1:], [])          # no spectra: everything is the data


def test_a_uv_cdf_is_kept_when_it_is_the_only_copy_of_that_detector():
    from services.workspace_loader import split_duplicate_exports
    files = [(n, b'') for n in ('run_MS1 +TIC SCAN ESI_spectra.cdf', 'run_MS1 +TIC SCAN ESI.cdf', 'run_VWD1A.cdf',
                                'run_VWD1B.cdf', 'run.dx_VWD1B.CSV')]
    kept, duplicates = split_duplicate_exports(files)
    assert [n for n, _ in kept] == ['run_MS1 +TIC SCAN ESI_spectra.cdf', 'run_VWD1A.cdf', 'run.dx_VWD1B.CSV']
    assert [n for n, _ in duplicates] == ['run_MS1 +TIC SCAN ESI.cdf', 'run_VWD1B.cdf']


def test_a_folder_of_runs_is_one_group_and_tab_named_after_it_with_one_run_shown(sequence):
    from callbacks.viz_callbacks import _get_overlay_group, _tab_id_from_group, _tab_label_from_group
    ws = Workspace()
    ws.open(str(sequence))
    runs = [b for b in ws._bundles() if b.bundle_type.value == 'lcms']
    keys = set()
    for bundle in runs:
        entry = ws.store.get_data_by_id(bundle.bundle_id)
        keys.add(_get_overlay_group(entry['metadata'], entry, file_id=bundle.bundle_id, dm=ws.store))
    assert len(keys) == 1 and _tab_label_from_group(keys.pop(), ws.store) == 'HalB'
    assert sorted(ws.store.visibility_state.get(b.bundle_id, True) for b in runs) == [False, True]


def test_pda_channels_and_spectra_are_read_from_the_3d_data():
    from services.data_store import Bundle, BundleType, Component, ComponentRole
    from src.ingestion.types import ColumnRole, SourceCategory
    import pandas as pd
    times = np.linspace(0, 3, 301)
    wavelengths = np.arange(190.0, 400.0, 1.2)
    peak = np.exp(-((times - 2.5) / 0.03) ** 2)[:, None]
    matrix = 800 * peak * np.exp(-((wavelengths - 205) / 12) ** 2)[None, :] + 300 * peak * \
        np.exp(-((wavelengths - 263) / 10) ** 2)[None, :]
    frame = pd.DataFrame({'Time (min)': times, 'Absorbance (mAU)': matrix[:, np.argmin(abs(wavelengths - 263))]})
    meta = {'wavelength_nm': 263, 'auto_x': 'Time (min)', 'auto_y': ['Absorbance (mAU)'],
            'pda': {'time_min': times, 'wavelength_nm': wavelengths, 'absorbance_mau': matrix}}
    component = Component('ala-263', ComponentRole.PRIMARY, 'ala.txt', dataframe=frame,
                          column_roles={'Time (min)': ColumnRole.X, 'Absorbance (mAU)': ColumnRole.Y},
                          source_category=SourceCategory.CHROMATOGRAM, metadata=meta)
    ws = Workspace()
    ws.store.add_bundle(Bundle('ala', 'ala 10 mM', BundleType.CHROMATOGRAM, component_list=[component]))
    spectrum = ws.uv_spectrum('ala 10 mM', 2.5)
    assert [m['nm'] for m in spectrum['maxima']] == [pytest.approx(205, abs=1.5), pytest.approx(263, abs=1.5)]
    added = ws.add_uv_channel(205)
    assert added[0]['measured_nm'] == pytest.approx(205, abs=0.7)
    rows = ws.detect_peaks(channel='205 nm')                        # the nearest recorded wavelength's channel
    assert [round(r['rt'], 2) for r in rows] == [2.5]


def test_agent_changes_are_in_the_history_with_the_agent_and_the_project_says_so(sequence, tmp_path):
    from services.research_package import agent_note, record_analysis_state
    ws = Workspace()
    ws.agent = {'via': 'MCP agent', 'client': 'claude-code', 'version': '2.1', 'model': 'Claude Opus 5.5'}
    ws.open(str(sequence))
    ws.integrate(8.2, 8.6, samples='4', channel='+TIC', name='m634 compound')
    steps = [e for e in ws.store.analysis_history if e.get('agent')]
    assert any('Integrated 8.2-8.6 min as m634 compound' in e['trigger'][0] for e in steps)
    assert 'claude-code 2.1, model Claude Opus 5.5' in agent_note(ws.store)
    again = Workspace()
    again.open(ws.save(str(tmp_path / 'p')))
    assert again.store.agent_clients == [ws.agent] and 'MCP agent' in agent_note(again.store)
    # Styling and view state are not analysis history; processing settings are.
    before = len(again.store.analysis_history)
    record_analysis_state(again.store, settings={'line_width': 3, 'font_family': 'Arial'})
    assert len(again.store.analysis_history) == before
    record_analysis_state(again.store, settings={'line_width': 3, 'smooth_window': 21})
    assert again.store.analysis_history[-1]['changes']['plot_settings']['after'] == {'smooth_window': 21}
    assert 'no MCP agent' in agent_note(Workspace().store)


def test_the_results_workbook_has_a_sheet_per_tab_with_areas_peaks_and_traces(sequence, tmp_path):
    from openpyxl import load_workbook
    ws = Workspace()
    ws.open(str(sequence))
    ws.integrate(8.2, 8.6, samples=['4', 'CONTROL 4'], channel='+TIC', name='m634 compound')
    book = load_workbook(ws.export_results(str(tmp_path / 'results')))
    assert book.sheetnames == ['Overview', 'HalB', 'HalB data', 'Samples', 'History']
    rows = list(book['HalB'].iter_rows(values_only=True))
    header = next(r for r in rows if r and r[0] == 'Compound')
    areas = rows[rows.index(header) + 1]
    assert header[:5] == ('Compound', 'Channel', 'RT (min)', '4', 'CONTROL 4')
    assert areas[0] == 'm634 compound' and areas[3] == pytest.approx(2 * areas[4], rel=0.05)
    peaks = [r for r in rows[rows.index(header) + 2:] if r and r[0] in ('4', 'CONTROL 4')]
    assert len(peaks) == 2
    data = list(book['HalB data'].iter_rows(min_row=4, max_row=6, values_only=True))
    assert data[0][0].startswith('4 · ') and data[1][:2] == ('Time (min)', 'Signal') and data[2][0] is not None


def test_an_eic_is_a_channel_with_an_editable_peak_that_survives_saving(sequence, tmp_path):
    ws = Workspace()
    ws.open(str(sequence))
    first = {r['sample']: r for r in ws.eic(ION, 0.5)}
    assert first['4']['channel'] == 'EIC 633.8 ±0.5 (+)' and first['4']['area'] > 0
    again = {r['sample']: r for r in ws.eic(ION, 0.5)}                  # replaced, not duplicated
    eic_peaks = [r for r in ws.peaks('4') if r['channel'].startswith('EIC')]
    assert len(eic_peaks) == 1 and eic_peaks[0]['area'] == pytest.approx(again['4']['area'])
    assert eic_peaks[0]['area'] == pytest.approx(2 * {r['sample']: r for r in ws.peaks('CONTROL 4')
                                                     if r['channel'].startswith('EIC')}['CONTROL 4']['area'], rel=0.02)
    bundle = ws._sample('4')
    assert bundle.primary.role.value != 'generic'                       # the run's primary channel is unchanged
    reopened = Workspace()
    reopened.open(ws.save(str(tmp_path / 'eic')))
    kept = [r for r in reopened.peaks('4') if r['channel'].startswith('EIC')]
    assert len(kept) == 1 and kept[0]['area'] == pytest.approx(eic_peaks[0]['area'])
    names = [c['name'] for c in next(x for x in reopened.samples() if x['name'] == '4')['channels']]
    assert 'EIC 633.8 ±0.5 (+)' in names


def test_removing_the_ms_channel_removes_its_eics(sequence):
    from services.channel_visibility import channel_controls
    ws = Workspace()
    ws.open(str(sequence))
    ws.add_eic_channel(ION, 0.5)
    bundle = ws._sample('4')
    keys = [c['key'] for c in channel_controls(ws.store, [bundle.bundle_id])]
    assert 'eic:+:633.8:0.5' in keys and keys.index('eic:+:633.8:0.5') > keys.index('ms:positive')
    ws.store.file_removal_service.remove_channels([bundle.bundle_id], 'ms:positive')
    left = [c['key'] for c in channel_controls(ws.store, [bundle.bundle_id])]
    assert 'ms:positive' not in left and not any(k.startswith('eic:') for k in left)


def test_live_mode_follows_the_agent_to_its_channel_and_view(sequence):
    """After an agent call the page shows what it worked on (services/live_focus): the sample and channel of the
    call, switched on when hidden, and the MS scan or the peak's mass spectrum."""
    from services.live_focus import focus
    ws = Workspace()
    ws.open(str(sequence))
    bundle = ws._sample('4')
    tic = ws._channel(bundle, '+TIC')
    ws.store.visibility_state[ws.store.peak_owner_id(tic.component_id)] = False     # the user shows UV only
    rows = ws.integrate(8.2, 8.6, samples='4', channel='+TIC', name='m634')
    view = focus(ws, 'integrate', [8.2, 8.6, '4', '+TIC', 'm634'], {}, rows)
    assert view['active'] == tic.component_id and view['results'] == 'peak-analysis'
    assert view['zoom'][0] < rows[0]['rt_start'] and view['zoom'][1] > rows[0]['rt_end']      # the plot zooms there
    assert ws.store.visibility_state[ws.store.peak_owner_id(tic.component_id)] is True
    view = focus(ws, 'mass_spectrum', ['4'], {'rt': 8.4}, ws.mass_spectrum('4', rt=8.4))
    assert view['side'] == 'ms-scan' and abs(TIMES[view['scan']] - 8.4) < 0.03
    view = focus(ws, 'mass_spectrum', ['4'], {'peak': 'm634'}, ws.mass_spectrum('4', peak='m634'))
    assert view['side'] == 'ms-peak-mass' and view['peak'] == rows[0]['peak_id']
    assert focus(ws, 'samples', [], {}, ws.samples()) is None                         # reads leave the view
    assert focus(ws, 'integrate', [1, 2], {}, {'error': 'x'}) is None


@pytest.mark.parametrize('nested', [False, True])
def test_data_folder_reports_projects_without_sending_them_to_ingestion(sequence, nested):
    folder = sorted(sequence.iterdir())[0] if nested else sequence
    projects = [folder / 'analysis.gecko', folder / 'fair.ZIP']
    for project in projects:
        project.write_bytes(b'not raw data')
    result = Workspace().open(str(sequence))
    assert result['not_readable'] == []
    assert result['projects'] == [str(p) for p in projects]
    assert 'open_data' in result['note']
    assert result['samples'] == 2


def test_folder_containing_only_a_project_leaves_the_workspace_empty(tmp_path):
    project = tmp_path / 'analysis.gecko'
    project.write_bytes(b'not raw data')
    result = Workspace().open(str(tmp_path))
    assert result['added'] == [] and result['not_readable'] == []
    assert result['projects'] == [str(project)]


def test_runs_sharing_a_sample_name_are_named_by_their_sample_id_in_one_group(tmp_path):
    """Labs put the compound in the Sample Name and tell the runs apart by Sample ID: such runs are named by their ID
    and grouped (one file-tree group and plot tab) under the name; a run without an ID leaves the set alone."""
    fixture = open(os.path.join(ROOT, 'tests', 'ingestion', 'fixtures', 'shimadzu_ascii.txt'), encoding='utf-8').read()
    folder = tmp_path / 'cal'
    folder.mkdir()
    for i, sample_id in enumerate(('15mM', '3.25mM')):
        text = fixture.replace('Synthetic PDA standard', 'Phenylethylamine').replace('fixture-001', sample_id)
        (folder / f'run_{i}.txt').write_text(text, encoding='utf-8')
    ws = Workspace()
    ws.open(str(folder))
    bundles = list(ws.store.iter_bundles())
    assert sorted(b.name for b in bundles) == ['15mM', '3.25mM']
    assert {ws.store.group_overrides[b.bundle_id][2] for b in bundles} == {'Phenylethylamine'}
    assert all(b.metadata['sample_name'] == 'Phenylethylamine' for b in bundles)        # the Metadata keeps it


def test_a_renamed_group_names_its_tab_and_tree_group_and_is_kept_in_the_project(tmp_path):
    from types import SimpleNamespace
    from callbacks.viz_callbacks import _tab_id_from_group, _tab_label_from_group
    from services import session_service
    from services.data_store import DataStore
    key = ('folder', 'Time (min)', 'Phenylethylamine')
    store = DataStore()
    assert _tab_label_from_group(key, store) == 'Phenylethylamine'
    store.group_labels = {_tab_id_from_group(key): 'PEA calibration 8 Oct'}
    assert _tab_label_from_group(key, store) == 'PEA calibration 8 Oct'
    restored = DataStore()
    session_service.restore_session(restored, session_service.build_session(store))
    assert _tab_label_from_group(key, restored) == 'PEA calibration 8 Oct'
    assert _tab_label_from_group(key, SimpleNamespace()) == 'Phenylethylamine'


def test_live_mode_opens_the_tab_of_the_call(sequence):
    from services.live_focus import focus
    ws = Workspace()
    ws.open(str(sequence))
    rows = ws.integrate(8.2, 8.6, channel='+TIC', name='m634')
    assert focus(ws, 'integrate', [8.2, 8.6], {'channel': '+TIC', 'name': 'm634'}, rows)['results'] == 'peak-analysis'
    done = ws.calibrate('m634', {'4': 2.0, 'CONTROL 4': 1.0}, channel='+TIC')
    view = focus(ws, 'calibrate', ['m634', {'4': 2.0, 'CONTROL 4': 1.0}], {'channel': '+TIC'}, done)
    model = next(iter(ws.store.calibration_service.get_all_models().values()))
    assert view == {'results': 'calibration', 'calibration': model.model_id}
    assert focus(ws, 'eic', [ION, 0.5], {}, ws.eic(ION, 0.5))['side'] == 'ms-eic'


def test_the_analysis_guides_answer_the_judgement_calls():
    from services.agent_guide import TOPICS, guide
    assert {'wavelength', 'integration', 'deconvolution', 'calibration', 'identification', 'conversion'} <= set(TOPICS)
    assert 'saturat' in guide('Wavelength')['guide'] and '15 %' in guide('calibration')['guide']
    assert 'Topics' in guide('nope')['error']


def test_a_project_saved_by_name_goes_to_the_projects_folder(sequence, tmp_path, monkeypatch):
    monkeypatch.setenv('GECKO_SETTINGS_DIR', str(tmp_path / 'settings'))
    ws = Workspace()
    ws.open(str(sequence))
    saved = ws.save('PEA calibration')
    assert saved == str(tmp_path / 'settings' / 'Projects' / 'PEA calibration.gecko') and os.path.exists(saved)



def test_the_agent_zoom_covers_every_peak_of_the_tab_with_little_baseline(sequence):
    from services.channel_visibility import visible_component_id
    from services.live_focus import peaks_range
    ws = Workspace()
    ws.open(str(sequence))
    bundle = next(b for b in ws.store.iter_bundles() if b.name == '4')
    tab, zoom = peaks_range(ws.store, visible_component_id(ws.store, bundle))
    a, b = zoom['xaxis.range']
    assert tab.startswith('tab-') and a < 8.4 < b and b - a < 2      # the shown TIC: its peak at 8.4 of 0-14 min


def test_a_peak_name_is_always_text():
    """An imported peak named by its time came back from the peak table as a number (and an empty cell as NaN): every
    later naming step then failed with a TypeError, also for the agent."""
    from services.peak_service import Peak, PeakService
    peak = Peak('p', 'f', 6.79, 6.79, 6.7, 6.9, 1.0, 1.0, 0.2)
    assert peak.name == '6.79'
    peak.label = float('nan')
    assert peak.label == ''
    assert PeakService.has_custom_name(peak)


def test_open_reads_only_the_matching_files_and_samples_can_be_removed(sequence):
    run = next(p for p in sequence.iterdir() if p.is_dir() and 'CONTROL' not in p.name)
    ws = Workspace()
    result = ws.open(str(run), match='*15-53-54*')
    assert result['added'] and len(ws.samples()) == 1
    other = Workspace()
    other.open(str(sequence))
    names = [s['name'] for s in other.samples()]
    remaining = other.remove_samples([names[0]])
    assert [s['name'] for s in remaining] == names[1:]


def test_after_the_agent_only_its_eic_is_shown(sequence):
    from services.channel_visibility import channel_identity
    from services.live_focus import focus
    ws = Workspace()
    ws.open(str(sequence))
    ws.add_eic_channel(ION, 0.5)
    ws.add_eic_channel(BACKGROUND, 0.5)
    rows = ws.eic(ION, 0.5)
    focus(ws, 'eic', [ION, 0.5], {}, rows)
    shown = {channel_identity(c)[0] for b in ws.store.iter_bundles() for c in b.component_list
             if channel_identity(c) and str(channel_identity(c)[0]).startswith('eic:')
             and ws.store.visibility_state.get(ws.store.peak_owner_id(c.component_id), True)}
    assert len(shown) == 1 and '633.8' in next(iter(shown))


def test_a_run_file_opened_on_its_own_later_is_a_sample_of_its_own_name(sequence, tmp_path):
    import shutil
    run = next(p for p in sequence.iterdir() if p.is_dir() and 'CONTROL' not in p.name)
    spectra = next(run.glob('*_spectra.cdf'))
    copy = tmp_path / 'again'
    copy.mkdir()
    shutil.copy(spectra, copy / spectra.name)
    ws = Workspace()
    ws.open(str(spectra))
    ws.open(str(copy / spectra.name))
    names = [s['name'] for s in ws.samples()]
    assert len(names) == 2 and len(set(names)) == 2 and names[1].endswith('(+TIC)')
